import docx
from docx.shared import Inches

doc = docx.Document()
doc.add_heading('Tabel 4.1 Hasil Pengujian Fungsional Sistem', level=1)

table_data = [
    ["NO", "Fitur yang Diuji", "Skenario Pengujian", "Hasil yang Diharapkan", "Hasil Aktual", "Status"],
    ["1", "Login Sistem", "Mengisi email dan password yang valid", "Sistem mengarahkan ke halaman Dashboard Filament", "Dashboard berhasil ditampilkan", "LULUS"],
    ["2", "Login Sistem", "Mengisi email atau password salah", "Sistem menolak akses dan menampilkan pesan error", "Pesan kesalahan kredensial muncul", "LULUS"],
    ["3", "Tambah Tiket Gangguan", "Membuat tiket keluhan baru dengan mengisi data pelanggan, alamat, dan keluhan secara lengkap", "Data tiket tersimpan dan muncul di tabel NamaPelanggans", "Data berhasil tersimpan dan tampil di tabel", "LULUS"],
    ["4", "Update Status Tiket", "Mengubah status tiket pada form edit (misal: dari Open ke In Progress)", "Status tiket diperbarui di database dan warna badge pada tabel berubah", "Status dan visualisasi badge berhasil di-update", "LULUS"],
    ["5", "Tautan Google Maps", "Mengklik tombol \"📍 Lihat Lokasi\" pada baris data tiket", "Sistem membuka tab peramban baru yang memuat Google Maps sesuai tautan koordinat", "Google Maps berhasil termuat di tab baru", "LULUS"],
    ["6", "Kelola Master Data", "Menambahkan atau menghapus entitas Keluhan/Status di Master Data", "Opsi dropdown keluhan/status pada form pembuatan tiket akan langsung menyesuaikan", "Perubahan pada Master Data langsung terefleksi", "LULUS"],
    ["7", "Manajemen User", "Admin menghapus kredensial teknisi/pegawai lain", "Akun teknisi terhapus dan tidak bisa login kembali", "Data user berhasil terhapus secara permanen", "LULUS"]
]

table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'

# Add header
hdr_cells = table.rows[0].cells
for i, col_name in enumerate(table_data[0]):
    hdr_cells[i].text = col_name

# Add rows
for row in table_data[1:]:
    row_cells = table.add_row().cells
    for i, cell_data in enumerate(row):
        row_cells[i].text = cell_data

doc.save('/root/perkuliahan/projek-magang/Tabel_4_1_Pengujian_Sistem.docx')
